import copy
from pathlib import Path
import numpy as np
import pytest
import yaml
from wingcheck.statics import cantilever,assemble,point_load,static_solve,segmented_grid
from wingcheck.dynamics import modes
from wingcheck.sections import component_section,section
from wingcheck.excel import Workbook,convert
from wingcheck.model import material_compliance,normalize
from wingcheck.postprocess import convergence,parse_dat


def uniform(EI=2000,mu=2):
    return lambda x:{'EI':EI,'mu':mu}


def test_cantilever_tip_force():
    x = np.linspace(0,2,1001)
    r = cantilever(x,np.zeros_like(x),np.full_like(x,2000),tip_force=100)
    assert r['w'][-1] == pytest.approx(100*2**3/(3*2000),rel=2e-6)
    assert r['root_force'] == pytest.approx(100)
    assert r['root_moment'] == pytest.approx(200)


def test_cantilever_uniform_load():
    x = np.linspace(0,2,2001)
    r = cantilever(x,np.full_like(x,50),np.full_like(x,2000))
    assert r['w'][-1] == pytest.approx(50*2**4/(8*2000),rel=2e-6)
    assert r['root_force'] == pytest.approx(100)
    assert r['root_moment'] == pytest.approx(100)


def test_supported_midpoint_and_overhang():
    x = np.linspace(-2,2,41)
    K,M,f = assemble(x,uniform())
    point_load(f,x,0,-100)
    fixed = [2*10,2*30]
    u,r = static_solve(K,f,fixed)
    assert u[40] == pytest.approx(-100*2**3/(48*2000),rel=1e-8)
    assert r[fixed] == pytest.approx([50,50],rel=1e-8)
    assert sum(r[fixed])+sum(f[::2]) == pytest.approx(0,abs=1e-7)
    assert u[0] > 0 and u[-2] > 0
    assert u[40] != 0 # root is not clamped


def test_uniform_cantilever_frequency():
    x = np.linspace(0,2,41)
    K,M,f = assemble(x,uniform())
    freq,_ = modes(K,M,[0,1],3)
    expected = 1.875104068711961**2/(2*np.pi*2**2)*np.sqrt(2000/2)
    assert freq[0] == pytest.approx(expected,rel=1e-5)


def test_variable_EI_convergence():
    answers = []
    # Exact unit-load solution integral F*(L-x)^2/EI(x) dx.
    from scipy.integrate import quad
    exact = quad(lambda x:100*(2-x)**2/(2000*(1+x)),0,2)[0]
    for n in [4,8,16]:
        x = np.linspace(0,2,n+1)
        K,M,f = assemble(x,lambda xx:{'EI':2000*(1+xx),'mu':2})
        point_load(f,x,2,100)
        u,_ = static_solve(K,f,[0,1])
        answers.append(abs(u[-2]-exact))
    assert answers[2] < answers[1] < answers[0]
    assert answers[2]/exact < 1e-5


def test_mass_and_central_mass_once():
    x = np.linspace(-2,2,41)
    K,M,f = assemble(x,uniform(mu=2),central_mass=5,central_inertia=3)
    translation = np.zeros(len(M))
    translation[::2] = 1
    assert translation@M@translation == pytest.approx(4*2+5)
    assert M[41,41] > 3
    _,M0,_ = assemble(x,uniform(mu=2))
    delta = M-M0
    assert np.sum(delta) == pytest.approx(8)


def test_elastic_centroid_and_steiner():
    p = [{'E':10,'area':2,'z':1,'I':.2,'density':3},
         {'E':20,'area':2,'z':-1,'I':.2,'density':3}]
    s = component_section(p)
    assert s['neutral_axis'] == pytest.approx(-1/3)
    assert s['EI'] == pytest.approx(10*(.2+2*(4/3)**2)+20*(.2+2*(2/3)**2))
    assert s['mu'] == 12


def test_orthotropic_admissibility_and_reciprocity():
    m = {'E1':135e9,'E2':8e9,'E3':8e9,'G12':4.5e9,'G13':4.5e9,'G23':3e9,
         'nu12':.28,'nu13':.28,'nu23':.35,'density':1550}
    S = material_compliance(m)
    assert np.all(S==S.T)
    assert -S[1,0]*m['E2'] == pytest.approx(m['nu12']*m['E2']/m['E1'])
    with pytest.raises(np.linalg.LinAlgError):
        material_compliance({**m,'nu23':3})
    with pytest.raises(ValueError):
        material_compliance({**m,'G13':None})


def test_units():
    assert convert(2800,'mm')/2 == pytest.approx(1.4)
    assert convert(135000,'MPa') == 135e9
    assert convert(1,'mm4') == 1e-12
    assert convert(64,'%') == .64
    with pytest.raises(ValueError):
        convert(1,'inch')


def test_xlsx_missing_formula_cache_and_label_mapping(tmp_path):
    import openpyxl
    p = tmp_path/'test.xlsx'
    b = openpyxl.Workbook()
    b.active['A1']='Spannweite'
    b.active['B1']=2800
    b.active['C1']='=B1/2'
    b.save(p)
    reader = Workbook(p)
    assert reader.resolve({'label':'Spannweite'})[0] == 2800
    assert reader.inventory()['warnings']
    with pytest.raises(ValueError):
        reader.get('Sheet','C1')


def test_actual_workbook_import_precedence_hash():
    cfg = yaml.safe_load(Path('config.yaml').read_text(encoding='utf-8'))
    cfg['analysis']['allow_provisional'] = False # Explicit strict mode, independent of user's active assumptions.
    b = Workbook(cfg['excel'])
    old_hash = b.sha256
    cfg['overrides']['total_lift'] = {'value':2000,'unit':'N'}
    m = normalize(b,cfg,{'total_lift':1500})
    assert m['parameters']['total_lift'] == 1500
    assert m['half_span'] == pytest.approx(1.4)
    assert len(m['stations']) >= 12 # Original 11 plus exact minimum-width/wall crossings.
    assert m['missing']
    assert len(next(p for p in m['provenance'] if p['parameter']=='total_lift')['overrides'])==2
    assert Workbook(cfg['excel']).sha256 == old_hash
    assert any(v.get('cached') for v in b.cells.values())


def test_provisional_model_does_not_claim_confirmed_inputs():
    cfg = yaml.safe_load(Path('config.yaml').read_text(encoding='utf-8'))
    cfg['analysis']['allow_provisional']=True
    cfg['material']['provisional']=True
    cfg['geometry']['manufacturing_plan_confirmed']=False
    cfg['flight']['load_kind']='comparison'
    m = normalize(Workbook(cfg['excel']),cfg)
    assert not m['missing']
    assert any('Fertigung NICHT bestätigt' in s for s in m['warnings'])
    assert any('Vergleichslast' in s for s in m['warnings'])


def test_discrete_plies_and_variant_stiffness():
    cfg = yaml.safe_load(Path('config.yaml').read_text(encoding='utf-8'))
    model = normalize(Workbook(cfg['excel']),cfg)
    cfg['geometry']['ply_thickness_m'] = .0001 # explicit synthetic verification parameter
    model['material']['density'] = 1550
    for x in [.199,.201,.399,.401,.799,.801,1.399]:
        a,b = section(model,cfg,x,'A'),section(model,cfg,x,'B')
        assert b['EI']>=a['EI']
        assert b['mu']>=a['mu']
    assert section(model,cfg,.199)['plies_top']==10
    assert section(model,cfg,.201)['plies_top']==9
    assert section(model,cfg,1.4)['plies_top']==2


def test_convergence_requires_three_levels():
    row = {'flight_tip_m':1.,'test_root_m':2.,'frequency_Hz':3.}
    assert not convergence([row,row],.02)['converged']
    assert convergence([row,row,row],.02)['converged']


def test_full_pipeline_synthetic_inputs():
    from wingcheck.analysis import compute
    cfg = yaml.safe_load(Path('config.yaml').read_text(encoding='utf-8'))
    cfg['material'].update(yaml.safe_load(Path('examples/material.example.yaml').read_text(encoding='utf-8'))['material'])
    cfg['geometry'].update(ply_thickness_m=.0001,manufacturing_plan_confirmed=True)
    cfg['flight']['load_kind']='limit'
    cfg['analysis']['elements_per_interval']=3
    model = normalize(Workbook(cfg['excel']),cfg)
    result = compute(model,cfg)
    assert result['status']=='ausgeführt'
    a,b = [r['summary'] for r in result['variants'].values()]
    assert b['half_spar_mass_kg'] >= a['half_spar_mass_kg']
    assert b['flight_tip_m'] <= a['flight_tip_m']
    assert abs(a['test_reaction_residual_N']) < 1e-4
    assert len(a['frequencies_Hz'])==6


def test_solver_parser(tmp_path):
    p = tmp_path/'sample.dat'
    p.write_text(' displacements (vx,vy,vz) for set ALL and time 1\n\n 1 1.0E-3 0.0 -2.0E-2\n\n forces (fx,fy,fz) for set ALL and time 1\n\n 1 0.0 0.0 10.0\n',encoding='ascii')
    d = parse_dat(p)
    assert d['displacements'][1][2] == -.02
    assert d['forces'][1][2] == 10


def test_real_calculix_frequency_header(tmp_path):
    p = tmp_path/'modal.dat'
    p.write_text(' MODE NO    EIGENVALUE                       FREQUENCY\n'
                 '                                     REAL PART            IMAGINARY PART\n'
                 '                           (RAD/TIME)      (CYCLES/TIME     (RAD/TIME)\n\n'
                 '      1   0.1498628E+08   0.3871212E+04   0.6161225E+03   0.0000000E+00\n'
                 '     P A R T I C I P A T I O N   F A C T O R S\n',encoding='ascii')
    assert parse_dat(p)['frequencies']==[616.1225]


def test_segmented_jump_integral():
    from wingcheck.statics import segmented_cantilever
    from scipy.integrate import quad
    def sec(x):
        return {'EI':1000 if x<1 else 2000,'mu':1}
    x = segmented_grid([0,1,2],4)
    r = segmented_cantilever(x,np.full_like(x,50.),sec)
    expected = quad(lambda s:50*(2-s)**3/(2*sec(s)['EI']),0,1)[0]+quad(lambda s:50*(2-s)**3/(2*sec(s)['EI']),1,2)[0]
    assert r['w'][-1]==pytest.approx(expected,rel=1e-12)
