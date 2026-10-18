"""Read-only inventory, formula cache audit and configurable provenance mapping.

BIFF FORMULA tokens are retained even where xlrd cannot decompile a formula.
No Excel, VBA or external links are executed. Formulas are never evaluated.
"""
import hashlib
import json
import math
import struct
from pathlib import Path

import xlrd
import openpyxl
import olefile


UNITS = {"m": (1., "m"), "mm": (.001, "m"), "mm2": (1e-6, "m²"),
         "mm4": (1e-12, "m⁴"), "Nmm": (.001, "Nm"), "N/mm": (1000., "N/m"),
         "MPa": (1e6, "Pa"), "GPa": (1e9, "Pa"), "Pa": (1., "Pa"),
         "kg": (1., "kg"), "kg/m3": (1., "kg/m³"), "N": (1., "N"),
         "%": (.01, "1"), "1": (1., "1"), "Hz": (1., "Hz"),
         "g/m2": (.001, "kg/m²")}


def convert(value, unit):
    if unit not in UNITS:
        raise ValueError(f"Unbekannte Einheit {unit!r}; keine stille Umrechnung")
    value = float(value)
    if not math.isfinite(value):
        raise ValueError("Nichtendlicher Eingabewert")
    return value * UNITS[unit][0]


def address(row, col):
    return f"{xlrd.formula.colname(col)}{row+1}"


class Workbook:
    def __init__(self, path):
        self.path = Path(path)
        self.sha256 = hashlib.sha256(self.path.read_bytes()).hexdigest()
        self.cells, self.sheets, self.warnings = {}, [], []
        if self.path.suffix.lower() == ".xls":
            self._xls()
        elif self.path.suffix.lower() == ".xlsx":
            self._xlsx()
        else:
            raise ValueError("Nur echte .xls/.xlsx werden unterstützt")

    def _xls(self):
        b = xlrd.open_workbook(str(self.path), formatting_info=True)
        for s in b.sheets():
            self.sheets.append({"name": s.name, "rows": s.nrows, "columns": s.ncols})
            for r in range(s.nrows):
                for c in range(s.ncols):
                    cell = s.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_EMPTY:
                        continue
                    note = s.cell_note_map.get((r, c))
                    self.cells[(s.name, address(r, c))] = {
                        "value": cell.value, "type": cell.ctype,
                        "comment": note.text if note else None,
                        "error": cell.ctype == xlrd.XL_CELL_ERROR}
        with olefile.OleFileIO(str(self.path)) as ole:
            stream = "Workbook" if ole.exists("Workbook") else "Book"
            data = ole.openstream(stream).read()
        # BOUNDSHEET offsets identify worksheet streams; chart BOFs are not sheets.
        offsets, p = {}, 0
        while p+4 <= len(data):
            op, n = struct.unpack_from("<HH", data, p)
            a = data[p+4:p+4+n]
            if op == 0x85 and len(a) >= 8 and a[5] == 0:
                offsets[struct.unpack_from("<I", a)[0]] = len(offsets)
            p += 4+n
        for start, i in offsets.items():
            if i >= len(b.sheets()):
                continue
            s, p = b.sheet_by_index(i), start
            shared = []
            while p+4 <= len(data):
                op, n = struct.unpack_from("<HH", data, p)
                a = data[p+4:p+4+n]
                p += 4+n
                if op == 0x0A:
                    break
                if op == 0x4BC and n >= 10:
                    r0,r1,c0,c1 = struct.unpack_from('<HHBB', a)
                    ln = struct.unpack_from('<H', a, 8)[0]
                    shared.append((r0,r1,c0,c1,a[10:10+ln]))
                if op != 6 or n < 22:
                    continue
                r, c = struct.unpack_from("<HH", a)
                ln = struct.unpack_from("<H", a, 20)[0]
                item = self.cells.setdefault((s.name, address(r, c)), {"value": None})
                item["formula_tokens_hex"] = a[22:22+ln].hex()
                try:
                    item["formula"] = xlrd.formula.decompile_formula(
                        b, a[22:], ln, xlrd.formula.FMLA_TYPE_CELL, browx=r, bcolx=c)
                except Exception as exc:
                    item["formula"] = None
                    item["formula_decode_error"] = str(exc)
                item["cached"] = True
            for r0,r1,c0,c1,tokens in shared:
                for r in range(r0,r1+1):
                    for c in range(c0,c1+1):
                        item = self.cells.get((s.name,address(r,c)))
                        if item and item.get('cached'):
                            item['shared_formula_tokens_hex'] = tokens.hex()
                            try:
                                item['formula'] = xlrd.formula.decompile_formula(
                                    b,tokens,len(tokens),xlrd.formula.FMLA_TYPE_SHARED,browx=r,bcolx=c)
                            except Exception as exc:
                                item['formula_decode_error'] = str(exc)

    def _xlsx(self):
        formulas = openpyxl.load_workbook(self.path, data_only=False, read_only=False,
                                         keep_links=False)
        values = openpyxl.load_workbook(self.path, data_only=True, read_only=True,
                                       keep_links=False)
        for s in formulas:
            self.sheets.append({"name": s.title, "rows": s.max_row, "columns": s.max_column})
            for row in s:
                for c in row:
                    if c.value is None:
                        continue
                    formula = c.value if c.data_type == "f" else None
                    self.cells[(s.title, c.coordinate)] = {
                        "value": values[s.title][c.coordinate].value,
                        "formula": formula, "cached": bool(formula),
                        "error": c.data_type == "e",
                        "comment": c.comment.text if c.comment else None}
        formulas.close()
        values.close()

    def get(self, sheet, cell):
        item = self.cells.get((sheet, cell), {})
        v = item.get("value")
        if v is None or v == "" or item.get("error"):
            raise ValueError(f"Fehlender/fehlerhafter Wert oder Formelcache: {sheet}!{cell}")
        return v

    def inventory(self):
        errors = []
        for (s, c), v in self.cells.items():
            if v.get("error") or (v.get("cached") and v.get("value") in (None, "")):
                errors.append(f"{s}!{c}: fehlender/fehlerhafter Formelcache")
        return {"path": str(self.path.resolve()), "sha256": self.sha256,
                "sheets": self.sheets, "warnings": self.warnings+errors,
                "formula_policy": "Gespeicherte Ergebnisse sind Caches, nicht neu berechnet. xlrd-Formeltext ist diagnostisch (Klammerdarstellung nicht verlässlich); BIFF-Token sind die verlustfreie Referenz.",
                "cells": [{"sheet": s, "cell": c, **v} for (s,c),v in self.cells.items()]}

    def resolve(self, spec):
        """Explicit address takes precedence; unique labels can supply generic mappings."""
        sheet, cell = spec.get("sheet"), spec.get("cell")
        if not cell:
            hits = [(s,c) for (s,c),v in self.cells.items()
                    if (not sheet or s == sheet) and
                    str(v.get("value", "")).strip().casefold() == spec["label"].casefold()]
            if len(hits) != 1:
                raise ValueError(f"Beschriftung nicht eindeutig: {spec.get('label')}: {hits}")
            sheet, label_cell = hits[0]
            r, c = openpyxl.utils.cell.coordinate_to_tuple(label_cell)
            dr, dc = spec.get("offset", [0, 1])
            cell = f"{openpyxl.utils.get_column_letter(c+dc)}{r+dr}"
        return self.get(sheet, cell), sheet, cell


def dump_json(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
